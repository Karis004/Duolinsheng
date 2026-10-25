import datetime
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import openpyxl

# Tests never load real reminder credentials or contact the reminder service.
sys.modules["SendKey"] = types.SimpleNamespace(user_name_list={})
import review
import storage


class PersistentRecordsTests(unittest.TestCase):
    def test_data_directory_environment_and_legacy_default(self):
        project = str(Path(__file__).resolve().parents[1])
        with tempfile.TemporaryDirectory() as folder:
            env = dict(os.environ, PYTHONPATH=project)
            env["DUOLINSHENG_DATA_DIR"] = folder
            command = [sys.executable, "-c",
                       "from storage import data_path; print(data_path('record-Example.xlsx').resolve())"]
            selected = subprocess.check_output(command, env=env, text=True).strip()
            self.assertEqual(Path(selected), Path(folder) / "record-Example.xlsx")
            env.pop("DUOLINSHENG_DATA_DIR")
            legacy = subprocess.check_output(command, env=env, cwd=folder, text=True).strip()
            self.assertEqual(Path(legacy), Path(folder) / "record-Example.xlsx")

    def test_learning_review_and_mistakes_use_persistent_files(self):
        with tempfile.TemporaryDirectory() as folder:
            data = Path(folder) / "data"
            data.mkdir()
            record = data / "record-Example.xlsx"
            today = datetime.date.today()
            yesterday = (today - datetime.timedelta(days=1)).isoformat()
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Content", "Description", "First", "Next", "Times",
                          "Mistakes", "Unused", "Days", "Last"])
            sheet.append(["alpha", "alpha meaning", None, None, None, 0, None, 4, yesterday])
            sheet.append(["bravo", "bravo meaning", yesterday, today.isoformat(), "1", 0])
            workbook.save(record)
            workbook.close()

            # A different working directory proves paths do not depend on /app.
            previous = Path.cwd()
            os.chdir(folder)
            try:
                decoy = Path("record-Example.xlsx")
                decoy.write_bytes(b"working-directory file must remain untouched")
                original = hashlib.sha256(decoy.read_bytes()).hexdigest()
                with patch.object(storage, "DATA_DIR", data):
                    words, days = review.review("Example", 1)
                    self.assertEqual(days, 5)
                    self.assertTrue((data / "today_words-Example.txt").is_file())
                    self.assertFalse(Path("today_words-Example.txt").exists())
                    # Existing files store the streak as text; keep that behavior.
                    self.assertEqual(review.get_review("Example"), (words, str(days)))

                    review.report_mistake(["alpha\n"], "Example")
                    errors, counts = review.error_book("Example")
                    self.assertEqual(errors, ["alpha\n", "alpha meaning\n"])
                    self.assertEqual(counts, [1])
                    with patch.object(review.SendKey, "user_name_list", {"Example": "test-only"}), \
                            patch.object(review.requests, "get") as request:
                        review.send_reminder()
                        request.assert_not_called()

                self.assertEqual(hashlib.sha256(decoy.read_bytes()).hexdigest(), original)
                workbook = openpyxl.load_workbook(record)
                self.assertEqual(workbook.active.cell(2, 9).value, today.isoformat())
                self.assertEqual(workbook.active.cell(2, 6).value, 1)
                workbook.close()
            finally:
                os.chdir(previous)


if __name__ == "__main__":
    unittest.main()
